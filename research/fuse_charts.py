# -*- coding: utf-8 -*-
"""
fuse_charts.py — 读取 research 下 *-*-* 各 eps 的 sweep_result.xlsx，
融合生成四类图表并输出到 research 目录：

  1) 01_侧铣点铣对比.png          侧铣 vs 点铣散点图（多 eps，颜色区分）
  2) 02_程序运行计时_分片数.png    程序运行计时 vs 分片数（多 eps 折线）
  3) 03_刀轨提速_分片数.png        刀轨加工提速比 vs 分片数（多 eps 折线）
  4) 04_误差与过切欠切_eps_*.png   每个 eps 一图，合并误差与过切/欠切统计

依赖: openpyxl, matplotlib（需中文字体，见 _setup_font）。
"""
import os
import glob
from pathlib import Path

import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.lines import Line2D

HERE = Path(__file__).resolve().parent

# 各 eps 的统一配色（跨图 1/2/3 保持一致）
EPS_PALETTE = ["#1f77b4", "#ff7f0e", "#2ca02c", "#d62728"]


def _setup_font():
    for _f in ["Microsoft YaHei", "SimHei", "SimSun"]:
        if any(f.name == _f for f in font_manager.fontManager.ttflist):
            plt.rcParams["font.sans-serif"] = [_f]
            break
    plt.rcParams["axes.unicode_minus"] = False


def load_sweep(xlsx_path):
    """读取单个 sweep_result.xlsx，返回 (eps, rows)。rows 为按 total 排序的 dict 列表。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    eps = None
    if "meta" in wb.sheetnames:
        ws_meta = wb["meta"]
        try:
            eps = float(ws_meta["B1"].value)
        except Exception:
            eps = None
    ws = wb["machining_sweep"]
    headers = None
    rows = []
    for row in ws.iter_rows(values_only=True):
        if headers is None:
            headers = list(row)
            continue
        d = dict(zip(headers, row))
        if d.get("total") is None:
            continue
        rows.append(d)
    wb.close()
    rows.sort(key=lambda r: r.get("total") or 0)
    return eps, rows


def collect_all():
    """扫描 research 下 *-*-* 目录，按 eps 升序返回 [(eps, rows), ...]。"""
    data = []
    for xlsx in sorted(glob.glob(str(HERE / "*-*-*" / "sweep_result.xlsx"))):
        eps, rows = load_sweep(xlsx)
        if eps is None or not rows:
            continue
        data.append((eps, rows))
    data.sort(key=lambda x: x[0])
    return data


def _eps_label(eps):
    return f"eps={eps:g}"


def chart1_scatter(data, out_path):
    """侧铣 vs 点铣散点图：x=平均误差, y=提速比，横/竖虚线及标注颜色对应 eps。"""
    eps_color = {eps: EPS_PALETTE[i] for i, (eps, _) in enumerate(data)}

    fig, ax = plt.subplots(figsize=(9, 5.5), dpi=140)
    handles = []
    for eps, rows in data:
        c = eps_color[eps]
        mean_err = [r["mean_error_mm"] for r in rows]
        speedup = [r["speedup"] for r in rows]
        ax.scatter(mean_err, speedup, color=c, s=22, alpha=0.85, zorder=3)
        handles.append(Line2D([], [], color=c, marker="o", linestyle="",
                              markersize=6, label=f"{_eps_label(eps)} 侧铣"))
        # 点铣基线：竖直虚线 x=eps，水平虚线 y=1，星形标记
        ax.axvline(eps, color=c, linestyle="--", linewidth=1.2, alpha=0.75, zorder=2)
        ax.axhline(1.0, color=c, linestyle="--", linewidth=1.2, alpha=0.75, zorder=2)
        ax.scatter([eps], [1.0], marker="*", s=260, color=c, zorder=4,
                   edgecolors="white", linewidths=0.6)
        ax.annotate(f"{eps:g}", xy=(eps, 1.0), xytext=(6, 10),
                    textcoords="offset points", color=c, fontsize=9, fontweight="bold",
                    zorder=5)
    # 点铣基线图例项（所有 eps 共享 speedup=1 水平线）
    handles.append(Line2D([], [], color="#555555", marker="*", linestyle="--",
                          markersize=11, label="点铣基线 (speedup=1, 残留=eps)"))

    ax.set_xscale("log")
    ax.set_xlabel("平均误差 Mean Error (mm)")
    ax.set_ylabel("提速比 Speedup")
    ax.set_title("侧铣(直纹面拟合) vs 点铣基线")
    ax.grid(True, alpha=0.3, which="both")
    ax.legend(handles=handles, fontsize=8, loc="best")
    fig.tight_layout()
    fig.savefig(out_path)
    plt.close(fig)
    print(f"[ok] {out_path.name}")


def chart2_runtime(data, out_path):
    """程序运行计时 vs 分片数：程序运行总耗时 = 拟合耗时 + 刀轨计算耗时。"""
    eps_color = {eps: EPS_PALETTE[i] for i, (eps, _) in enumerate(data)}

    fig, ax = plt.subplots(figsize=(9, 5.5), dpi=140)
    for eps, rows in data:
        c = eps_color[eps]
        totals = [r["total"] for r in rows]
        runtime = [(r["fit_time_s"] or 0.0) + (r["toolpath_time_s"] or 0.0)
                   for r in rows]
        ax.plot(totals, runtime, color=c, linewidth=1.6, marker="o", markersize=3.5,
                label=_eps_label(eps))
    ax.set_xlabel("分片数 Total Slices")
    ax.set_ylabel("程序运行计时 (s)")
    ax.set_title("程序运行计时 vs 分片数")
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=9, title="eps")
    fig.tight_layout()
    fig.savefig(out_path)
    plt.close(fig)
    print(f"[ok] {out_path.name}")


def chart3_speedup(data, out_path):
    """刀轨加工提速比 vs 分片数。"""
    eps_color = {eps: EPS_PALETTE[i] for i, (eps, _) in enumerate(data)}

    fig, ax = plt.subplots(figsize=(9, 5.5), dpi=140)
    for eps, rows in data:
        c = eps_color[eps]
        totals = [r["total"] for r in rows]
        speedup = [r["speedup"] for r in rows]
        ax.plot(totals, speedup, color=c, linewidth=1.6, marker="o", markersize=3.5,
                label=_eps_label(eps))
    ax.set_xlabel("分片数 Total Slices")
    ax.set_ylabel("刀轨加工提速比 Speedup")
    ax.set_title("刀轨加工提速比 vs 分片数")
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=9, title="eps")
    fig.tight_layout()
    fig.savefig(out_path)
    plt.close(fig)
    print(f"[ok] {out_path.name}")


def chart4_error_overcut(eps, rows, out_path):
    """单个 eps：合并误差统计与过切/欠切统计（6 条曲线，图例颜色区分）。"""
    totals = [r["total"] for r in rows]
    series = [
        ("min_error_mm", "最小误差 min", "#1f77b4", "-"),
        ("max_error_mm", "最大误差 max", "#ff7f0e", "-"),
        ("mean_error_mm", "平均误差 mean", "#2ca02c", "-"),
        ("rms_error_mm", "RMS 误差 rms", "#9467bd", "-"),
        ("max_overcut_mm", "最大过切 over-cut", "#d62728", "--"),
        ("max_undercut_mm", "最大欠切 under-cut", "#17becf", "--"),
    ]
    fig, ax = plt.subplots(figsize=(9, 5.5), dpi=140)
    for key, label, color, ls in series:
        ys = [r.get(key) for r in rows]
        ax.plot(totals, ys, color=color, linestyle=ls, linewidth=1.5,
                marker="o", markersize=3, label=label)
    ax.set_xlabel("分片数 Total Slices")
    ax.set_ylabel("误差 (mm)")
    ax.set_title(f"误差与过切/欠切统计 (eps={eps:g})")
    ax.grid(True, alpha=0.3)
    ax.legend(fontsize=8, ncol=2)
    fig.tight_layout()
    fig.savefig(out_path)
    plt.close(fig)
    print(f"[ok] {out_path.name}")


def main():
    _setup_font()
    data = collect_all()
    if not data:
        print("[error] 未找到 *-*-*/sweep_result.xlsx 数据文件")
        return
    print(f"读取到 {len(data)} 组 eps: {[f'{eps:g}' for eps, _ in data]}")

    chart1_scatter(data, HERE / "01_侧铣点铣对比.png")
    chart2_runtime(data, HERE / "02_程序运行计时_分片数.png")
    chart3_speedup(data, HERE / "03_刀轨提速_分片数.png")
    for eps, rows in data:
        fname = f"04_误差与过切欠切_eps_{eps:g}.png"
        chart4_error_overcut(eps, rows, HERE / fname)
    print("完成。")


if __name__ == "__main__":
    main()
